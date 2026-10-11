from pathlib import Path
import re, json
from zipfile import ZipFile
from lxml import etree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

BASE = Path(__file__).resolve().parent / 'kinetic-garden'
OUT = BASE / 'Living_Spatial_Infrastructure_Prologue_and_606_v0.3.docx'
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.78)
sec.left_margin = sec.right_margin = Inches(.82)
sec.header_distance = sec.footer_distance = Inches(.35)
CONTENT = 6.86

for style in doc.styles:
    if style.type in (1, 2):
        style.font.name = 'Malgun Gothic'
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style.element.get_or_add_rPr()
        fonts = rpr.find(qn('w:rFonts'))
        if fonts is None:
            fonts = OxmlElement('w:rFonts'); rpr.append(fonts)
        for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn('w:' + attr), 'Malgun Gothic')
        for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme', 'csTheme'):
            fonts.attrib.pop(qn('w:' + attr), None)
        col = rpr.find(qn('w:color'))
        if col is not None:
            for attr in ('themeColor', 'themeTint', 'themeShade'):
                col.attrib.pop(qn('w:' + attr), None)
        pb = style.element.find('.//' + qn('w:pBdr'))
        if pb is not None: pb.getparent().remove(pb)
        if style.type == 1:
            ppr = style.element.get_or_add_pPr()
            snap = OxmlElement('w:snapToGrid'); snap.set(qn('w:val'), '0'); ppr.append(snap)

grid = sec._sectPr.find(qn('w:docGrid'))
if grid is not None: sec._sectPr.remove(grid)

normal = doc.styles['Normal']
normal.font.size = Pt(11)
normal.paragraph_format.line_spacing = 1.45
normal.paragraph_format.space_after = Pt(7)
normal.paragraph_format.widow_control = True
for name, size, before, after in [('Title', 27, 0, 12), ('Subtitle', 13, 0, 10), ('Heading 1', 19, 17, 9), ('Heading 2', 14, 14, 7), ('Heading 3', 12, 11, 5)]:
    st = doc.styles[name]
    st.font.size = Pt(size)
    st.font.bold = name != 'Subtitle'
    st.font.italic = False
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True
    st.paragraph_format.line_spacing = 1.2

def inline(p, text):
    pattern = r'(\[[^\]]+\]\([^\)]+\)|\*\*[^*]+\*\*|`[^`]+`)'
    for token in re.split(pattern, text):
        if not token: continue
        m = re.fullmatch(r'\[([^\]]+)\]\(([^\)]+)\)', token)
        if m:
            link = OxmlElement('w:hyperlink')
            rid = p.part.relate_to(m[2], RT.HYPERLINK, is_external=True)
            link.set(qn('r:id'), rid)
            run = OxmlElement('w:r'); pr = OxmlElement('w:rPr')
            col = OxmlElement('w:color'); col.set(qn('w:val'), '225A73'); pr.append(col)
            u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); pr.append(u)
            run.append(pr); t = OxmlElement('w:t'); t.text = m[1]; run.append(t)
            link.append(run); p._p.append(link)
        else:
            r = p.add_run(token[2:-2] if token.startswith('**') else token.strip('`'))
            if token.startswith('**'): r.bold = True

def para(text, style=None):
    p = doc.add_paragraph(style=style); inline(p, text); return p

def clean_heading(text):
    text = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text).replace('**', '')
    return re.sub(r'[^\w\s가-힣]', ' ', text).strip()

def table(rows):
    n = len(rows[0])
    first = rows[0][0]
    # Narrative case records stay in prose so readers can follow facts and inference.
    if first == '사례':
        for row in rows[1:]:
            doc.add_heading(clean_heading(row[0]), level=3)
            for label, value in zip(['확인한 구현', '이번 기획에 연결하는 해석', '적용 경계'], row[1:]):
                para('**' + label + '**  ' + value)
        return
    # User decision records become an editable response area.
    if first in ('내가 결정할 항목', '선택할 항목'):
        for row in rows[1:]:
            p = para('**' + row[0] + '**')
            p.paragraph_format.keep_with_next = True
            para('선택 후보  ' + row[1])
            p = para('내 선택과 이유  ')
            p.paragraph_format.space_after = Pt(14)
        return
    if first == 'ID': weights = [.10, .45, .45]
    elif n == 2: weights = [.30, .70]
    elif n == 3: weights = [.23, .41, .36]
    elif first == '후보': weights = [.27, .28, .34, .11]
    else: weights = [.18, .28, .29, .25]
    t = doc.add_table(rows=1, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, w in zip(t.columns, weights): col.width = Inches(CONTENT*w)
    pr = t._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b = OxmlElement('w:' + edge)
        b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4'); b.set(qn('w:color'), 'D9D9D9')
        borders.append(b)
    pr.append(borders)
    for i, values in enumerate(rows):
        row = t.rows[0] if i == 0 else t.add_row()
        trpr = row._tr.get_or_add_trPr()
        if i == 0: trpr.append(OxmlElement('w:tblHeader'))
        trpr.append(OxmlElement('w:cantSplit'))
        for j, (cell, value) in enumerate(zip(row.cells, values)):
            cell.width = Inches(CONTENT*weights[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cp = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for edge, twips in [('top', 95), ('bottom', 95), ('left', 100), ('right', 100)]:
                e = OxmlElement('w:' + edge); e.set(qn('w:w'), str(twips)); e.set(qn('w:type'), 'dxa'); margins.append(e)
            cp.append(margins)
            if i == 0:
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E7EEF2'); cp.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.line_spacing = 1.3
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = False
            inline(p, value)
            for r in p.runs:
                r.font.size = Pt(10.5)
                if i == 0: r.bold = True
    doc.add_paragraph().paragraph_format.space_after = Pt(3)

def markdown(text):
    lines = text.splitlines(); i = 0
    while i < len(lines):
        line = lines[i].strip(); i += 1
        if not line or line == '---': continue
        if line.startswith('|'):
            rows = [[v.strip() for v in line.strip('|').split('|')]]
            while i < len(lines) and lines[i].strip().startswith('|'):
                row = [v.strip() for v in lines[i].strip().strip('|').split('|')]; i += 1
                if all(re.fullmatch(r':?-+:?', v) for v in row): continue
                rows.append(row)
            table(rows); continue
        m = re.match(r'^(#{1,4})\s+(.*)', line)
        if m:
            if len(m[1]) == 1: continue
            doc.add_heading(clean_heading(m[2]), level=min(len(m[1]), 3))
            continue
        if re.match(r'^\d+\. ', line):
            para(re.sub(r'^\d+\. ', '', line), 'List Number'); continue
        if line.startswith('- '): para(line[2:], 'List Bullet'); continue
        para(line)

doc.core_properties.title = '살아있는 공간 인프라 제안서 프롤로그와 비초에 606 사례 분석'
doc.core_properties.subject = '주거와 식생의 통합 공간 및 설치와 장기 사용을 위한 설계 연구'
doc.core_properties.author = 'Kinetic Home'
doc.core_properties.version = '0.3'
doc.add_paragraph('살아있는 공간 인프라 제안서', 'Title')
doc.add_paragraph('프롤로그와 비초에 606 아키텍처 사례 분석', 'Subtitle')
para('Living Spatial Infrastructure  ·  Kinetic Home  ·  2026년 10월 11일  ·  v0.3')
para('이 제안서는 어머니의 거실정원을 정리하기 어려웠던 경험에서 출발해, 인간의 생활과 식물의 생육을 함께 설계하는 공간 인프라를 탐구한다. 앞부분은 실제 공간 사례로 질문을 확장하고, 뒷부분은 비초에 606의 설치와 장기 사용 원리를 정원 시스템에 연결한다.')
para('핵심 제안은 오래 유지할 기반과 바꾸어 쓸 생활·식재 구성을 함께 계획하는 것이다. 606의 조합이 건축처럼 느껴진다는 사용자의 직감을 분석의 출발점으로 삼고, 추가 후보 5개를 통해 구조·돌봄·변경 권한·물리적 접점을 비교한다. 형태와 구동 방식, 적용 범위와 후보 채택은 사용자가 결정할 수 있도록 남겨 둔다.')
para('본문의 확인은 공개 원문, 해석은 사례를 연결한 판단, 제안과 가설은 아직 시험하지 않은 가능성을 뜻한다. 이번 DOCX는 개념 프롤로그와 아키텍처 사례 분석을 묶은 문서다. 국가별 환경과 시장·제품 상세 조사는 별도 종합 보고서 REPORT.md와 연구 보드에 이어진다.')
doc.add_heading('문서의 구성', level=2)
for label in ['1 식물이 공간을 형성하는 프롤로그', '2 설치하고 오래 쓰는 시스템과 비초에 606', '3 추가 후보 5개와 물리적 사용자 인터페이스', '4 장기 사용의 조건과 다음 선택']:
    para(label)

doc.add_heading('1 식물이 공간을 형성하는 프롤로그', level=1)
prologue = (BASE/'LIVING_SPATIAL_INFRASTRUCTURE.md').read_text(encoding='utf-8')
prologue = '\n'.join(prologue.splitlines()[4:])
prologue = prologue.replace('해당 국가의 주거 환경 비교는 뒤의 기존 국가별 조사에 유지했다.', '해당 국가의 주거 환경 비교는 별도 종합 보고서의 국가별 조사에 정리했다.')
prologue = prologue.replace('뒤의 레일·모듈러 A~D', '종합 보고서의 레일·모듈러 A~D')
prologue = prologue.replace('아래의 거실·베란다 환경', '별도 종합 보고서의 거실·베란다 환경')
markdown(prologue)
doc.add_heading('2 설치하고 오래 쓰는 공간 인프라', level=1)
cases = (BASE/'606_ARCHITECTURE_REFERENCES.md').read_text(encoding='utf-8')
lines = cases.splitlines()[5:]
for i, line in enumerate(lines):
    if line.startswith('## 추가 후보 1 '):
        lines.insert(i, '#IGNORE\n'); break
cases = '\n'.join(lines).replace('#IGNORE\n', '##SECTION3\n')
# Add part headings at natural transitions without changing the source prose.
for chunk in re.split(r'(##SECTION3|## 장기 사용을 설계 조건으로 옮기는 질문)', cases):
    if chunk == '##SECTION3': doc.add_heading('3 추가 후보 5개와 선택 지도', level=1)
    elif chunk == '## 장기 사용을 설계 조건으로 옮기는 질문':
        doc.add_heading('4 장기 사용의 조건과 다음 선택', level=1)
        doc.add_heading('장기 사용을 설계 조건으로 옮기는 질문', level=2)
    else: markdown(chunk)

fp = sec.footer.paragraphs[0]
fp.alignment = 2
r = fp.add_run('Kinetic Home   '); r.font.size = Pt(9)
field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); fp._p.append(field)
doc.save(OUT)
with ZipFile(OUT) as z:
    xml = etree.fromstring(z.read('word/document.xml'))
    ns = {'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
    txt = ''.join(xml.xpath('//w:t/text()', namespaces=ns))
    assert all(s in txt for s in ['606', 'USM', 'Habraken', 'Bosch', 'GARDENA', 'Neutrik', 'MFO', '미정'])
    assert not xml.xpath('//w:trHeight[@w:hRule="exact"]', namespaces=ns)
    report = {'path':str(OUT), 'paragraphs':len(doc.paragraphs), 'tables':len(doc.tables), 'hyperlinks':len(xml.xpath('//w:hyperlink',namespaces=ns)), 'characters':len(txt), 'size':OUT.stat().st_size, 'renderStatus':'pending'}
    (BASE/'DOCX_VERIFICATION.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
